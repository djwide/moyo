"""Turn an uploaded note, PDF, Word file, or web page into plain text."""

from __future__ import annotations

from io import BytesIO

MAX_DOCUMENT_BYTES = 4_000_000
MAX_NOTE_CHARS = 400_000


def document_to_text(raw: bytes, suffix: str) -> str:
    """Recover UTF-8 text. Empty and image-only files raise ValueError."""
    if not raw or len(raw) > MAX_DOCUMENT_BYTES:
        raise ValueError("Document is empty or exceeds 4 MB.")
    ext = suffix.lower()
    if not ext.startswith("."):
        ext = f".{ext}"
    if ext in {".md", ".txt"}:
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ValueError("That note is not valid UTF-8.") from exc
    elif ext == ".pdf":
        text = _pdf_text(raw)
    elif ext == ".docx":
        text = _docx_text(raw)
    elif ext in {".html", ".htm"}:
        text = _html_text(raw)
    else:
        raise ValueError(f"Unsupported document type: {suffix}")
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        raise ValueError("No text could be recovered from that document.")
    if len(text) > MAX_NOTE_CHARS:
        raise ValueError("Recovered text exceeds the note limit.")
    return text


def _pdf_text(raw: bytes) -> str:
    try:
        from pdfminer.high_level import extract_text
    except ImportError as exc:
        raise ValueError("PDF text extraction is unavailable.") from exc
    try:
        return extract_text(BytesIO(raw)) or ""
    except Exception as exc:
        raise ValueError("No text could be recovered from that document.") from exc


def _docx_text(raw: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ValueError("Word text extraction is unavailable.") from exc
    try:
        document = docx.Document(BytesIO(raw))
    except Exception as exc:
        raise ValueError("No text could be recovered from that document.") from exc
    parts: list[str] = []
    for paragraph in document.paragraphs:
        line = paragraph.text.strip()
        if line:
            parts.append(line)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


_HTML_DROP_TAGS = ("script", "style", "noscript", "template", "svg", "iframe", "canvas", "object", "embed")
_HTML_BLOCK_TAGS = (
    "p",
    "div",
    "section",
    "article",
    "li",
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "tr",
    "blockquote",
    "pre",
    "figcaption",
    "header",
    "footer",
    "nav",
    "main",
    "aside",
    "table",
    "ul",
    "ol",
    "br",
)


def _html_text(raw: bytes) -> str:
    """Visible reading text from a saved web page, without markup or scripts."""
    try:
        from bs4 import BeautifulSoup
        from bs4.element import Comment
    except ImportError as exc:
        raise ValueError("HTML text extraction is unavailable.") from exc
    try:
        soup = BeautifulSoup(raw, "lxml")
    except Exception:
        try:
            soup = BeautifulSoup(raw, "html.parser")
        except Exception as exc:
            raise ValueError("No text could be recovered from that document.") from exc
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    for tag in list(soup.find_all(_HTML_DROP_TAGS)):
        tag.decompose()
    for comment in list(soup.find_all(string=lambda value: isinstance(value, Comment))):
        comment.extract()
    for tag in list(soup.find_all(True)):
        if tag.parent is None:
            continue
        if _html_tag_hidden(tag):
            tag.decompose()
    for image in list(soup.find_all("img")):
        alt = " ".join(str(image.get("alt") or "").split())
        image.replace_with(f" {alt} " if alt else "")
    for line_break in list(soup.find_all("br")):
        line_break.replace_with("\n")
    root = _html_reading_root(soup)
    for tag in list(root.find_all(_HTML_BLOCK_TAGS)):
        tag.insert_before("\n")
        tag.insert_after("\n")
    lines = [" ".join(line.split()) for line in root.get_text("", strip=False).splitlines()]
    body = "\n".join(line for line in lines if line)
    if title and not body.lower().startswith(title.lower()):
        return f"{title}\n\n{body}".strip()
    return body.strip()


def _html_tag_hidden(tag) -> bool:
    if tag.has_attr("hidden"):
        return True
    if str(tag.get("aria-hidden") or "").strip().lower() == "true":
        return True
    style = "".join(str(tag.get("style") or "").split()).lower()
    return "display:none" in style or "visibility:hidden" in style


def _html_reading_root(soup):
    """Prefer the page's main article when one is present."""
    main = soup.find("main")
    if main is not None and len(main.get_text(" ", strip=True)) >= 80:
        return main
    articles = [
        tag
        for tag in soup.find_all("article")
        if tag.find_parent("article") is None and len(tag.get_text(" ", strip=True)) >= 80
    ]
    if len(articles) == 1:
        return articles[0]
    if len(articles) > 1:
        return _HtmlFragment(articles)
    return soup.body or soup


class _HtmlFragment:
    """Several article nodes, read in document order."""

    def __init__(self, tags: list) -> None:
        self._tags = tags

    def find_all(self, names):
        found = []
        for tag in self._tags:
            if tag.name in names:
                found.append(tag)
            found.extend(tag.find_all(names))
        return found

    def get_text(self, separator: str = "", strip: bool = False) -> str:
        return "\n".join(tag.get_text(separator, strip=strip) for tag in self._tags)
